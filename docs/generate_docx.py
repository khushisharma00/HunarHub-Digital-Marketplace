import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def generate_docx():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    md_file = os.path.join(script_dir, 'project_report.md')
    docx_file = os.path.join(script_dir, 'project_report.docx')

    doc = Document()

    # Configure Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    with open(md_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("HUNARHUB")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(234, 88, 12) # Terracotta orange

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Digital Marketplace for Local Micro-Entrepreneurs\nB.Tech Project Report")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(14)
    sub_run.font.bold = True
    sub_run.font.color.rgb = RGBColor(71, 85, 105)

    author_p = doc.add_paragraph()
    author_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author_run = author_p.add_run("Author: Khushi Sharma | Date: October 2026\nDepartment of Computer Science & Engineering")
    author_run.font.name = "Arial"
    author_run.font.size = Pt(10)
    author_run.font.italic = True
    author_run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    in_table = False
    table_lines = []

    for line in lines:
        stripped = line.strip()

        # Handle Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines.append(stripped)
            continue
        elif table_lines:
            # Process buffered table
            filtered = [r for r in table_lines if not re.match(r'^\|[\s\-:|]+\|$', r)]
            if filtered:
                rows_data = [[c.strip() for c in r.strip('|').split('|')] for r in filtered]
                num_cols = len(rows_data[0])
                table = doc.add_table(rows=len(rows_data), cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.style = 'Table Grid'

                for row_idx, row in enumerate(rows_data):
                    for col_idx, cell_value in enumerate(row):
                        if col_idx < num_cols:
                            cell = table.cell(row_idx, col_idx)
                            cell.text = cell_value
                            p = cell.paragraphs[0]
                            p.runs[0].font.name = "Arial"
                            p.runs[0].font.size = Pt(9)
                            if row_idx == 0:
                                p.runs[0].font.bold = True
                                set_cell_background(cell, "F1F5F9")
            table_lines = []

        if not stripped:
            continue

        if stripped.startswith('# '):
            continue # already handled title
        elif stripped.startswith('## '):
            h = doc.add_heading(level=1)
            run = h.add_run(stripped[3:])
            run.font.name = "Arial"
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
        elif stripped.startswith('### '):
            h = doc.add_heading(level=2)
            run = h.add_run(stripped[4:])
            run.font.name = "Arial"
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(2)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            text = stripped[2:]
            # Bold matching
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Arial"
                r.font.size = Pt(10)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            parts = re.split(r'(\*\*.*?\*\*)', stripped)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Arial"
                r.font.size = Pt(10)

    doc.save(docx_file)
    print(f"✅ Generated DOCX successfully at: {docx_file}")

if __name__ == '__main__':
    generate_docx()
